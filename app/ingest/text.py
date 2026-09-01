"""Plain-text extraction for indexing / search.

Best-effort: returns ``""`` when a type is unsupported or parsing fails.
"""

from __future__ import annotations

import csv
import io
import zipfile
from html.parser import HTMLParser
from pathlib import Path

MAX_CHARS = 2_000_000


class _Stripper(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._skip = False
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in ("script", "style"):
            self._skip = True

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style"):
            self._skip = False

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def _html(data: bytes) -> str:
    p = _Stripper()
    p.feed(data.decode("utf-8", "replace"))
    return " ".join("".join(p.parts).split())


def _plain(data: bytes) -> str:
    return data.decode("utf-8", "replace")


def _csv(data: bytes) -> str:
    rows = csv.reader(io.StringIO(data.decode("utf-8", "replace")))
    return "\n".join(" ".join(r) for r in rows)


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    out: list[str] = []
    for page in reader.pages:
        out.append(page.extract_text() or "")
        if sum(len(s) for s in out) > MAX_CHARS:
            break
    return "\n".join(out)


def _docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs if p.text)


def _xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out.append(" ".join("" if c is None else str(c) for c in row))
            if sum(len(s) for s in out) > MAX_CHARS:
                wb.close()
                return "\n".join(out)
    wb.close()
    return "\n".join(out)


def _ooxml_generic(path: Path) -> str:
    """Fallback for pptx / unknown OOXML: pull text nodes from the xml parts."""
    import re

    chunks: list[str] = []
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if name.endswith(".xml") and ("slide" in name or "document" in name or "sheet" in name):
                xml = zf.read(name).decode("utf-8", "replace")
                chunks += re.findall(r"<a:t>([^<]+)</a:t>", xml)
    return " ".join(chunks)


def extract_text(path: Path, mime: str, ext: str) -> str:
    try:
        if mime == "application/pdf":
            text = _pdf(path)
        elif ext == "docx" or mime.endswith("wordprocessingml.document"):
            text = _docx(path)
        elif ext in ("xlsx", "xlsm") or mime.endswith("spreadsheetml.sheet"):
            text = _xlsx(path)
        elif ext == "pptx" or mime.endswith("presentationml.presentation"):
            text = _ooxml_generic(path)
        elif mime in ("text/html", "application/xhtml+xml") or ext in ("html", "htm"):
            text = _html(path.read_bytes())
        elif mime == "text/csv" or ext in ("csv", "tsv"):
            text = _csv(path.read_bytes())
        elif mime.startswith("text/") or ext in ("txt", "md", "log", "json", "xml"):
            text = _plain(path.read_bytes())
        else:
            return ""
    except Exception:
        return ""
    return text[:MAX_CHARS].strip()
